import re
from pathlib import Path

from docx import Document


def md_to_docx(md_text: str) -> Document:
    doc = Document()

    # Simple Markdown to DOCX: handle headings, paragraphs, and list items.
    # This is intentionally minimal to avoid complex dependencies.
    lines = md_text.splitlines()
    in_list = False
    for line in lines:
        if not line.strip():
            doc.add_paragraph("")
            continue

        # Headings: lines starting with ### or ##
        m3 = re.match(r"^###\s+(.*)$", line)
        m2 = re.match(r"^##\s+(.*)$", line)
        if m3:
            p = doc.add_paragraph()
            run = p.add_run(m3.group(1).strip())
            run.bold = True
            continue
        if m2:
            p = doc.add_paragraph()
            run = p.add_run(m2.group(1).strip())
            run.bold = True
            continue

        # Lists: lines starting with "- "
        if line.lstrip().startswith("- "):
            text = line.lstrip()[2:]
            doc.add_paragraph(text, style='List Bullet')
            in_list = True
            continue
        else:
            in_list = False

        # Inline citation brackets like [1] left as-is.
        doc.add_paragraph(line)

    return doc


def main():
    src = Path("/workspace/llm-security-intro-related-methodology.md")
    dst = Path("/workspace/llm-security-intro-related-methodology.docx")
    if not src.exists():
        raise SystemExit(f"Source not found: {src}")
    md_text = src.read_text(encoding="utf-8")
    doc = md_to_docx(md_text)
    doc.save(dst)
    print(f"Wrote {dst}")


if __name__ == "__main__":
    main()

