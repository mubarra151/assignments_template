from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor


def add_title_slide(prs: Presentation, title: str, subtitle: str) -> None:
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = title
    slide.placeholders[1].text = subtitle


def add_bullets_slide(prs: Presentation, title: str, bullets: list[str]) -> None:
    slide_layout = prs.slide_layouts[1]  # Title and Content
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = title
    tf = slide.shapes.placeholders[1].text_frame
    tf.clear()
    for i, line in enumerate(bullets):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = line
        p.level = 0


def add_two_column_slide(prs: Presentation, title: str, left_lines: list[str], right_lines: list[str]) -> None:
    slide_layout = prs.slide_layouts[5]  # Title Only
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = title

    left = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(4.5), Inches(5))
    right = slide.shapes.add_textbox(Inches(5.0), Inches(1.5), Inches(4.5), Inches(5))

    ltf = left.text_frame
    ltf.clear()
    for i, text in enumerate(left_lines):
        p = ltf.paragraphs[0] if i == 0 else ltf.add_paragraph()
        p.text = text
        p.level = 0

    rtf = right.text_frame
    rtf.clear()
    for i, text in enumerate(right_lines):
        p = rtf.paragraphs[0] if i == 0 else rtf.add_paragraph()
        p.text = text
        p.level = 0


def build_presentation() -> Presentation:
    prs = Presentation()

    # Title
    add_title_slide(
        prs,
        title="Prompt Injection and Data Leakage in LLM Applications",
        subtitle="Thesis Mid-Evaluation"
    )

    # Agenda
    add_bullets_slide(
        prs,
        title="Agenda",
        bullets=[
            "Background and Motivation",
            "Threat Models",
            "Systems Under Test",
            "Attacks and Defenses",
            "Methodology and Metrics",
            "Plan and Risks",
            "References",
        ],
    )

    # Background and Motivation
    add_bullets_slide(
        prs,
        title="Background and Motivation",
        bullets=[
            "LLMs integrated in RAG and agentic systems",
            "Indirect prompt injection via untrusted content",
            "Training-time memorization and runtime data exfiltration",
            "Transferability of adversarial prompts across models",
            "Need for defense-in-depth with utility awareness",
        ],
    )

    # Threat Models
    add_bullets_slide(
        prs,
        title="Threat Models",
        bullets=[
            "Prompt injection (indirect): override, prompt exposure, unsafe tool use",
            "Data leakage: training-time memorization and runtime exfiltration",
            "Adversary: black-box model; controls untrusted content; no privileged tools",
        ],
    )

    # Systems Under Test
    add_bullets_slide(
        prs,
        title="Systems Under Test",
        bullets=[
            "RAG pipeline: retriever, ranker, prompt constructor, generator",
            "Tool-using agent: web fetch, sandboxed file read, simple API",
            "Models: one open-weight and one hosted (where permitted)",
        ],
    )

    # Attacks
    add_bullets_slide(
        prs,
        title="Attacks",
        bullets=[
            "Indirect prompt injection in documents/web pages",
            "Web/HTML payloads in comments/attributes/alt-text",
            "Universal jailbreak baselines for comparison",
            "Leakage: canary seeding for runtime and training-time",
        ],
    )

    # Defenses
    add_two_column_slide(
        prs,
        title="Defenses (Layered)",
        left_lines=[
            "Prompt/tool hardening",
            "  • Least-privilege prompts",
            "  • Instruction boundaries and refusals",
            "  • JSON-only function calls",
            "  • Parameter allowlists",
        ],
        right_lines=[
            "Content/Guards/Isolation",
            "  • HTML/JS/CSS neutralization",
            "  • Input/output guards; DLP; canaries",
            "  • Tool sandboxing; scoped creds",
            "  • Provenance logging",
        ],
    )

    # Methodology and Metrics
    add_bullets_slide(
        prs,
        title="Methodology and Metrics",
        bullets=[
            "Automated attack harness; multi-turn; paraphrase-robust",
            "Defense evaluation matrix; incremental layering; ablations",
            "Metrics: ASR, leak rate, tokens/time-to-leak, TPR/FPR, utility, latency/cost",
            "CIs over 3+ seeds; paraphrase robustness",
        ],
    )

    # Plan and Risks
    add_bullets_slide(
        prs,
        title="Plan and Risks",
        bullets=[
            "Finalize corpus and canary design",
            "Run baseline → add defenses → ablations",
            "Risks: compute limits, eval variance, adaptive attacks",
            "Mitigation: sample-efficient eval, fixed seeds, reproducible harness",
        ],
    )

    # References
    add_bullets_slide(
        prs,
        title="References",
        bullets=[
            "Greshake et al., AISec'23 / arXiv:2302.12173",
            "Carlini et al., USENIX Sec'21 / arXiv:2012.07805",
            "Carlini et al., arXiv:2202.07646",
            "Zou et al., arXiv:2307.15043",
            "Anthropic, arXiv:2401.05566",
        ],
    )

    return prs


def main() -> None:
    prs = build_presentation()
    prs.save("/workspace/thesis-presentation.pptx")
    print("Wrote /workspace/thesis-presentation.pptx")


if __name__ == "__main__":
    main()

